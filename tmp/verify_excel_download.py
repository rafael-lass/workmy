from collections import Counter, defaultdict
from decimal import Decimal
import json
from pathlib import Path
import openpyxl

path = Path('tmp/workmy-export-total.xlsx')
book = openpyxl.load_workbook(path, read_only=True, data_only=True)
assert book.sheetnames == ['Dashboard', 'Pagamentos', 'Clientes', 'Serviços', 'Projetos']
totals, states, ids = defaultdict(Decimal), Counter(), set()
for row in book['Pagamentos'].iter_rows(min_row=2, values_only=True):
    assert row[0] not in ids
    ids.add(row[0])
    states[row[7]] += 1
    totals[row[7]] += Decimal(str(row[11]))
assert len(ids) == 200000
for state, cell in [('PAGO','B7'),('PENDENTE','B8'),('CANCELADO','B9'),('ESTORNADO','B10')]:
    assert totals[state] == Decimal(str(book['Dashboard'][cell].value))
counts = {name: book[name].max_row - 1 for name in book.sheetnames[1:]}
assert counts == {'Pagamentos': 200000, 'Clientes': 10000, 'Serviços': 53, 'Projetos': 12000}
book.close()
report = {'counts': counts, 'payment_statuses': dict(states), 'totals_brl': {k: str(v) for k,v in totals.items()}, 'dashboard_reconciled': True, 'unique_payment_ids': len(ids)}
Path('docs/evidencias/exportacao-excel-2026-10-08/download-validado.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
