from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import pytest

from src.application.dto.excel_export import ExportPeriod
from src.infrastructure.documents import dashboard_excel as module


@pytest.fixture
def renderer(tmp_path, monkeypatch):
    mkstemp, temporary_directory = module.mkstemp, module.TemporaryDirectory
    monkeypatch.setattr(module, 'mkstemp', lambda **kw: mkstemp(dir=tmp_path, **kw))
    monkeypatch.setattr(module, 'TemporaryDirectory', lambda **kw: temporary_directory(dir=tmp_path, **kw))
    return module.DashboardExcelRenderer(workers=2, batch_size=2)


def read_package(path):
    result = {}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            book = openpyxl.load_workbook(BytesIO(archive.read(name)))
            try:
                sheet = book.active
                result[name] = (book.sheetnames, sheet.freeze_panes, sheet.auto_filter.ref,
                                [[(c.value, c.data_type, c.number_format) for c in row] for row in sheet])
            finally:
                book.close()
    return result


def records():
    return {
        'pagamentos': [dict(id=i, projeto_id=1, status='PENDENTE', valor=Decimal('123.45'),
                           data_vencimento=date(2026, 9, 30), data_pagamento=None,
                           criado_em=datetime(2026, 9, 1, 15, 30),
                           observacao='=SUM(1,2) & <teste> 😀 _x0041_\r\n Próxima linha') for i in range(1, 8)],
        'clientes': [dict(id=1, nome='  João & Cia  ', criado_em=date(2020, 1, 1))],
        'servicos': [dict(id=1, nome='Site', descricao='Descrição', descricao_documento=None)],
        'projetos': [dict(id=1, status='IN_PROGRESS', valor=Decimal('50'), recorrencia_ativa=True)],
    }


def test_parallel_chunks_equal_sequential_cells_styles_order_and_text(renderer, tmp_path):
    period = ExportPeriod.create('total', date(2026, 10, 8))
    outputs = []
    try:
        for mode in ('sequential', 'parallel'):
            path, timings, sizes = renderer.render(records(), period, mode)
            outputs.append(path)
            assert all(value >= 0 for value in timings.values())
            assert sizes['pacote_zip'] > 0
        sequential, parallel = [read_package(path) for path in outputs]
        assert sequential == parallel
        payments = parallel['workmy-pagamentos-total-2026-10-08.xlsx']
        assert payments[:3] == (['Pagamentos'], 'A2', 'A1:Q8')
        rows = payments[3]
        assert [row[0][0] for row in rows[1:]] == list(range(1, 8))
        assert rows[1][9][0] == datetime(2026, 9, 30)
        assert rows[1][11][0] == 123.45
        assert rows[1][12][0] == 'Sim'
        assert rows[1][13][1] == 's'
        legacy = tmp_path / 'reference.xlsx'
        renderer._render_one('Pagamentos', module.PAYMENT_COLUMNS, records()['pagamentos'], legacy)
        outputs.append(legacy)
        reference = openpyxl.load_workbook(legacy)
        try:
            # Openpyxl preserves Excel's inline string escape tokens; match XlsxWriter.
            assert rows[1][13][0] == reference.active['N2'].value
        finally:
            reference.close()
    finally:
        for path in outputs:
            path.unlink(missing_ok=True)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize('mode', ['sequential', 'parallel'])
def test_empty_and_worker_failure_cleanup(renderer, tmp_path, mode):
    period = ExportPeriod.create('3', date(2026, 10, 8))
    path, _, _ = renderer.render(dict.fromkeys(('pagamentos', 'clientes', 'servicos', 'projetos'), []), period, mode)
    try:
        assert all(len(book[3]) == 1 for book in read_package(path).values())
    finally:
        path.unlink()
    data = records()
    data['pagamentos'][2]['valor'] = float('nan')
    with pytest.raises(ValueError, match='Valor de célula'):
        renderer.render(data, period, mode)
    assert not list(tmp_path.iterdir())
