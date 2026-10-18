from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile
from decimal import Decimal
import httpx
import openpyxl
import pytest
import pytest_asyncio
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from src.application.dto.excel_export import ExportPeriod
from src.infrastructure.persistence.session import Base, get_db_session
from src.infrastructure.persistence.models import UsuarioModel, ClienteModel, ServicoModel, ProjetoModel, PagamentoModel
from src.infrastructure.security.jwt_service import create_access_token
from src.presentation.main import app
from src.presentation.rest import dashboard
from src.infrastructure.documents import dashboard_excel


@pytest.fixture
def fixed_today(monkeypatch):
    monkeypatch.setattr(dashboard, 'business_today', lambda: date(2027, 1, 20))


@pytest_asyncio.fixture
async def api(fixed_today, monkeypatch, tmp_path):
    original_mkstemp = dashboard_excel.mkstemp
    original_temporary_directory = dashboard_excel.TemporaryDirectory
    monkeypatch.setattr(dashboard_excel, 'mkstemp', lambda **kwargs: original_mkstemp(dir=tmp_path, **kwargs))
    monkeypatch.setattr(dashboard_excel, 'TemporaryDirectory',
                        lambda **kwargs: original_temporary_directory(dir=tmp_path, **kwargs))
    engine = create_async_engine('sqlite+aiosqlite:///:memory:')
    sessions = async_sessionmaker(engine)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with sessions.begin() as s:
        s.add_all([UsuarioModel(id=i, username=f'u{i}', email=f'u{i}@example.com', password_hash='test') for i in (1, 2, 3)])
        await s.flush()
        s.add_all([ClienteModel(id=i, usuario_id=1 if i < 3 else 2, nome='=1+1' if i == 1 else f'c{i}') for i in (1, 2, 3)])
        s.add_all([ServicoModel(id=i, usuario_id=1 if i < 3 else 2, nome=f's{i}') for i in (1, 2, 3)])
        await s.flush()
        s.add_all([ProjetoModel(id=i, usuario_id=1 if i < 3 else 2, cliente_id=i, servico_id=i, status='ARCHIVED', deletado_em=datetime(2027,1,1) if i == 1 else None) for i in (1,2,3)])
        await s.flush()
        specs = [(1,date(2026,10,31),'PAGO',25), (1,date(2026,11,1),'PAGO',100),
                 (1,date(2027,1,20),'PENDENTE',50), (1,date(2027,1,21),'PENDENTE',70),
                 (1,date(2026,12,1),'ESTORNADO',20), (1,date(2026,12,2),'CANCELADO',30),
                 (3,date(2027,1,1),'PAGO',999)]
        for project, due, status, value in specs:
            s.add(PagamentoModel(projeto_id=project, data_vencimento=due, status=status, valor=Decimal(value),
                data_pagamento=due if status in ('PAGO','ESTORNADO') else None,
                observacao='=HYPERLINK("http://example.com")', motivo_ajuste='Teste' if status in ('CANCELADO','ESTORNADO') else None))
    async def session_override():
        async with sessions() as s:
            yield s
    previous = app.dependency_overrides.get(get_db_session)
    app.dependency_overrides[get_db_session] = session_override
    try:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app), base_url='http://test') as client:
            yield client
    finally:
        assert not list(tmp_path.glob('*.zip')), 'O pacote temporário deve ser removido após o download'
        assert not list(tmp_path.glob('workmy-excel-*')), 'Os workbooks temporários devem ser removidos'
        if previous:
            app.dependency_overrides[get_db_session] = previous
        else:
            app.dependency_overrides.pop(get_db_session, None)
        await engine.dispose()


def auth(user=1):
    return {'Authorization': f'Bearer {create_access_token(user, f"u{user}", f"u{user}@example.com")}'}


@pytest.mark.parametrize('key,expected', [('3',date(2026,11,1)), ('6',date(2026,8,1)), ('12',date(2026,2,1)), ('total',None)])
def test_period_calendar_and_year_boundary(key, expected):
    assert ExportPeriod.create(key, date(2027,1,20)).start == expected


@pytest.mark.asyncio
async def test_export_auth_period_and_isolation(api):
    assert (await api.get('/api/dashboard/exportar.zip')).status_code == 401
    assert (await api.get('/api/dashboard/exportar.zip?periodo=2', headers=auth())).status_code == 422
    response = await api.get('/api/dashboard/exportar.zip?periodo=3', headers=auth())
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert response.headers['content-type'] == 'application/zip'
    assert 'attachment;' in response.headers['content-disposition']
    with ZipFile(BytesIO(response.content)) as archive:
        assert archive.namelist() == ['workmy-pagamentos-3-2027-01-20.xlsx', 'workmy-clientes-3-2027-01-20.xlsx',
            'workmy-servicos-3-2027-01-20.xlsx', 'workmy-projetos-3-2027-01-20.xlsx']
        assert len(archive.namelist()) == 4
        def workbook(key):
            filename = next(name for name in archive.namelist() if f'-{key}-' in name)
            return openpyxl.load_workbook(BytesIO(archive.read(filename)))
        payments, clients, services, projects = (workbook(key) for key in ('pagamentos','clientes','servicos','projetos'))
        assert payments.sheetnames == ['Pagamentos']
        assert clients.sheetnames == ['Clientes']
        assert services.sheetnames == ['Serviços']
        assert projects.sheetnames == ['Projetos']
        assert payments['Pagamentos'].max_row == 5
        assert clients['Clientes'].max_row == services['Serviços'].max_row == projects['Projetos'].max_row == 2
        assert clients['Clientes']['B2'].value == '=1+1' and clients['Clientes']['B2'].data_type == 's'
        assert payments['Pagamentos']['N2'].data_type == 's'
        assert payments['Pagamentos']['L2'].value == 100
        assert payments['Pagamentos']['J2'].is_date
        assert payments['Pagamentos'].freeze_panes == 'A2'
        assert payments['Pagamentos'].auto_filter.ref
        for book in (payments, clients, services, projects):
            book.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('period', ['6', '12'])
async def test_other_periods_include_older_payment(api, period):
    response = await api.get(f'/api/dashboard/exportar.zip?periodo={period}', headers=auth())
    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as archive:
        name = next(name for name in archive.namelist() if 'pagamentos' in name)
        book = openpyxl.load_workbook(BytesIO(archive.read(name)))
        assert book['Pagamentos'].max_row == 6
        assert sum(row[0] is not None for row in book['Pagamentos'].iter_rows(min_row=2, values_only=True)) == 5
        book.close()


@pytest.mark.asyncio
async def test_total_includes_unrelated_catalog_and_empty_user(api):
    total = await api.get('/api/dashboard/exportar.zip?periodo=total', headers=auth())
    with ZipFile(BytesIO(total.content)) as archive:
        books = {key: openpyxl.load_workbook(BytesIO(archive.read(next(n for n in archive.namelist() if key in n))))
                 for key in ('pagamentos','clientes','servicos','projetos')}
        assert books['pagamentos']['Pagamentos'].max_row == 7
        assert books['clientes']['Clientes'].max_row == books['servicos']['Serviços'].max_row == books['projetos']['Projetos'].max_row == 3
        [book.close() for book in books.values()]
    empty = await api.get('/api/dashboard/exportar.zip?periodo=total', headers=auth(3))
    with ZipFile(BytesIO(empty.content)) as archive:
        name = next(name for name in archive.namelist() if 'pagamentos' in name)
        book = openpyxl.load_workbook(BytesIO(archive.read(name)))
        assert book['Pagamentos'].max_row == 1
        book.close()
